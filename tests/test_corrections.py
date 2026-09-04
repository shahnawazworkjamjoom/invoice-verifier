from pathlib import Path
from tempfile import TemporaryDirectory
import gc
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from invoice_verifier.corrections import ACTION_HEADER, AMOUNT_HEADER, export_decision_workbook
from invoice_verifier.models import InvoiceRecord, VerificationResult


def make_result(decision="APPROVE", corrected_amount=None) -> VerificationResult:
    record = InvoiceRecord(
        excel_row=2, supplier="BARAKAT", brand="Subway", location="Test",
        order_number="PO-1", invoice_number="INV-1", currency="AED",
        po_amount=100, invoice_date="01-Sep-2026", attachment_url="https://example.invalid/a.pdf",
        unique_reference="x", record_id="1", payment_status="FULL", amount_to_pay=105,
        received_qty=1, tax_code="UAE_VAT_INC_5%")
    return VerificationResult(record, decision, "Internal detail must not be exported",
                              corrected_amount_to_pay=corrected_amount)


class DecisionExportTests(unittest.TestCase):
    def test_only_approval_cell_changes(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "source.xlsx"
            output = Path(folder) / "result.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            headers = ["Supplier Name", ACTION_HEADER, "Approver Remarks", "Record ID", AMOUNT_HEADER]
            values = ["BARAKAT", "", "existing remark", "1", 240.63]
            sheet.append(headers)
            sheet.append(values)
            workbook.save(source)
            workbook.close()

            count = export_decision_workbook(source, output, {2: make_result()})
            self.assertEqual(1, count)
            original_book = load_workbook(source, data_only=False, read_only=True)
            result_book = load_workbook(output, data_only=False, read_only=True)
            original = list(original_book.active.iter_rows(values_only=True))
            result = list(result_book.active.iter_rows(values_only=True))
            self.assertEqual(original[0], result[0])
            self.assertEqual("APPROVE", result[1][1])
            self.assertEqual(original[1][0], result[1][0])
            self.assertEqual(original[1][2:], result[1][2:])
            with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output) as result_zip:
                self.assertEqual(source_zip.namelist(), result_zip.namelist())
                changed_parts = {
                    name for name in source_zip.namelist()
                    if source_zip.read(name) != result_zip.read(name)
                }
                self.assertEqual({"xl/worksheets/sheet1.xml"}, changed_parts)
            original_book.close()
            result_book.close()
            del original_book, result_book
            gc.collect()

    def test_updates_existing_amount_cell_when_corrected(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "source.xlsx"
            output = Path(folder) / "result.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Supplier Name", ACTION_HEADER, AMOUNT_HEADER])
            sheet.append(["BARAKAT", "", 240.63])
            workbook.save(source)
            workbook.close()

            export_decision_workbook(source, output, {2: make_result(corrected_amount=252.68)})
            result_book = load_workbook(output, data_only=False, read_only=True)
            result = list(result_book.active.iter_rows(values_only=True))
            self.assertEqual("APPROVE", result[1][1])
            self.assertEqual(252.68, result[1][2])
            result_book.close()


if __name__ == "__main__":
    unittest.main()
