from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import VerificationResult


APP_DIR = Path(__file__).resolve().parent.parent
HISTORY_DIR = APP_DIR / "reports" / "history"
INDEX_NAME = "_index.csv"
AUDIT_SHEET = "Verification_Audit"

AUDIT_HEADERS = [
    "Excel Row",
    "Supplier",
    "Location",
    "Order No.",
    "Invoice No.",
    "Amount To Pay",
    "Currency",
    "Decision",
    "Remarks",
    "Final Amount Due (OCR)",
    "Finance Amount To Pay",
    "OCR Method",
    "Invoice Vendor",
    "Amount Evidence",
    "Manual Override",
    "Record ID",
    "Unique Reference",
    "Source File",
    "Audited At",
]


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return str(value)


def new_audit_path(source_path: str | Path, run_started_at: datetime | None = None) -> Path:
    run_started_at = run_started_at or datetime.now()
    stem = Path(source_path).stem
    # Keep only filesystem-safe characters so the file is findable after years.
    safe_stem = "".join(c if (c.isalnum() or c in ("-", "_")) else "_" for c in stem)[:80] or "workbook"
    filename = f"{safe_stem}_{run_started_at:%Y%m%d_%H%M%S}_audit.json"
    return HISTORY_DIR / filename


def result_to_dict(result: VerificationResult, audited_at: str, source_name: str) -> dict[str, Any]:
    record = result.record
    ocr = result.ocr_data or {}
    return {
        "excel_row": record.excel_row,
        "supplier": record.supplier,
        "brand": record.brand,
        "location": record.location,
        "order_number": record.order_number,
        "invoice_number": record.invoice_number,
        "currency": record.currency,
        "po_amount": _json_safe(record.po_amount),
        "invoice_date": record.invoice_date,
        "attachment_url": record.attachment_url,
        "unique_reference": record.unique_reference,
        "record_id": record.record_id,
        "payment_status": record.payment_status,
        "amount_to_pay": _json_safe(record.amount_to_pay),
        "received_qty": _json_safe(record.received_qty),
        "tax_code": record.tax_code,
        "decision": result.decision,
        "remarks": result.remarks,
        "checks": _json_safe(dict(result.checks or {})),
        "manual_override": (result.checks or {}).get("Manual Override", ""),
        "ocr_summary": _json_safe({
            "amount_total": ocr.get("amount_total"),
            "vendor_name": ocr.get("vendor_name"),
            "amount_evidence": ocr.get("amount_evidence"),
            "ocr_method": ocr.get("ocr_method"),
        }),
        "source_file": source_name,
        "audited_at": audited_at,
    }


def save_audit_snapshot(
    audit_path: str | Path,
    source_path: str | Path,
    results: dict[int, VerificationResult],
    run_started_at: datetime | None = None,
) -> Path:
    """Write the full verification story to JSON so it can be read after years.

    Called after verification finishes, after every manual APPROVE/DECLINE
    change, and after export. Rewriting the same file keeps one audit per run.
    """
    audit_path = Path(audit_path)
    source_path = Path(source_path)
    now = datetime.now()
    audited_at = now.isoformat(timespec="seconds")
    run_iso = (run_started_at or now).isoformat(timespec="seconds")

    ordered = [results[key] for key in sorted(results)]
    payload = {
        "app": "KOJ Invoice Verifier",
        "purpose": "Long-term audit: why the system took each APPROVE/DECLINE decision.",
        "source_name": source_path.name,
        "source_path": str(source_path),
        "run_started_at": run_iso,
        "saved_at": audited_at,
        "total": len(ordered),
        "approve": sum(1 for r in ordered if r.decision == "APPROVE"),
        "decline": sum(1 for r in ordered if r.decision == "DECLINE"),
        "results": [result_to_dict(r, audited_at, source_path.name) for r in ordered],
    }

    audit_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = audit_path.with_suffix(".tmp")
    tmp_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp_path.replace(audit_path)
    return audit_path


def log_history_index(audit_path: str | Path, source_path: str | Path,
                      results: dict[int, VerificationResult],
                      run_started_at: datetime | None = None) -> Path:
    """Append one line to reports/history/_index.csv so old runs are findable."""
    HISTORY_DIR.mkdir(parents=True, exist_ok=True)
    index_path = HISTORY_DIR / INDEX_NAME
    approve = sum(1 for r in results.values() if r.decision == "APPROVE")
    decline = sum(1 for r in results.values() if r.decision == "DECLINE")
    row = {
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "run_started_at": (run_started_at or datetime.now()).isoformat(timespec="seconds"),
        "source_name": Path(source_path).name,
        "audit_file": Path(audit_path).name,
        "total": len(results),
        "approve": approve,
        "decline": decline,
    }
    exists = index_path.exists()
    with index_path.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row.keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(row)
    return index_path


def embed_audit_sheet(exported_xlsx: str | Path, source_path: str | Path,
                      results: dict[int, VerificationResult]) -> int:
    """Add a Verification_Audit sheet to the exported workbook.

    The exported file then carries its own why-explanation, readable even in
    10 years without the JSON history folder.
    """
    exported_xlsx = Path(exported_xlsx)
    audited_at = datetime.now().isoformat(timespec="seconds")
    source_name = Path(source_path).name

    workbook = load_workbook(exported_xlsx)
    try:
        if AUDIT_SHEET in workbook.sheetnames:
            del workbook[AUDIT_SHEET]
        sheet = workbook.create_sheet(AUDIT_SHEET)
        sheet.append(AUDIT_HEADERS)
        for excel_row in sorted(results):
            result = results[excel_row]
            record = result.record
            checks = result.checks or {}
            sheet.append([
                record.excel_row,
                record.supplier,
                record.location,
                record.order_number,
                record.invoice_number,
                record.amount_to_pay,
                record.currency,
                result.decision,
                result.remarks,
                checks.get("Final Amount Due", ""),
                checks.get("Finance Amount To Pay", ""),
                checks.get("Local OCR", ""),
                checks.get("Invoice Vendor", ""),
                checks.get("Amount Evidence", ""),
                checks.get("Manual Override", ""),
                record.record_id,
                record.unique_reference,
                source_name,
                audited_at,
            ])
        widths = [11, 14, 30, 18, 20, 15, 9, 10, 60, 20, 20, 32, 20, 32, 32, 14, 18, 28, 20]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(exported_xlsx)
    finally:
        workbook.close()
    return len(results)
