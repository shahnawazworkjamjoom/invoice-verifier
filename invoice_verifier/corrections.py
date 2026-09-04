from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import VerificationResult
from .workbook import _normalise_header


ACTION_HEADER = "(APPROVE / DECLINE) Approver Action"
AMOUNT_HEADER = "Finance Amount To Pay"
SHARED_STRINGS_PATH = "xl/sharedStrings.xml"
SHEET_PATH = "xl/worksheets/sheet1.xml"


def _shared_string_indexes(xml: str) -> dict[str, int]:
    values: dict[str, int] = {}
    for index, block in enumerate(re.findall(r"<si(?:\s[^>]*)?>.*?</si>", xml, re.S)):
        parts = re.findall(r"<t(?:\s[^>]*)?>(.*?)</t>", block, re.S)
        text = "".join(re.sub(r"<[^>]+>", "", part) for part in parts)
        values[text] = index
    return values


def _ensure_decision_strings(xml: str) -> tuple[str, dict[str, int]]:
    indexes = _shared_string_indexes(xml)
    additions = [value for value in ("APPROVE", "DECLINE") if value not in indexes]
    if not additions:
        return xml, indexes
    next_index = len(re.findall(r"<si(?:\s[^>]*)?>", xml))
    blocks = []
    for value in additions:
        indexes[value] = next_index
        next_index += 1
        blocks.append(f"<si><t>{value}</t></si>")
    xml = xml.replace("</sst>", "".join(blocks) + "</sst>")
    unique = re.search(r'\buniqueCount="(\d+)"', xml)
    if unique:
        new_count = int(unique.group(1)) + len(additions)
        xml = xml[:unique.start(1)] + str(new_count) + xml[unique.end(1):]
    return xml, indexes


def _target_columns(source: Path) -> tuple[str, str]:
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        indexed = {_normalise_header(value): index + 1 for index, value in enumerate(header)}
        action_column = indexed.get(ACTION_HEADER)
        amount_column = indexed.get(AMOUNT_HEADER)
        if action_column is None:
            raise ValueError(f"Missing required column: {ACTION_HEADER}")
        if amount_column is None:
            raise ValueError(f"Missing required column: {AMOUNT_HEADER}")
        return get_column_letter(action_column), get_column_letter(amount_column)
    finally:
        workbook.close()


def export_decision_workbook(source: str | Path, destination: str | Path,
                             results: dict[int, VerificationResult]) -> int:
    """Copy the XLSX and change approval plus corrected existing amount cells only."""
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("Choose a new filename; the uploaded workbook cannot be overwritten.")
    action_column, amount_column = _target_columns(source)
    destination.parent.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(source, "r") as input_zip:
        has_shared_strings = SHARED_STRINGS_PATH in input_zip.namelist()
        shared_xml = None
        indexes = {}
        if has_shared_strings:
            shared_xml = input_zip.read(SHARED_STRINGS_PATH).decode("utf-8")
            shared_xml, indexes = _ensure_decision_strings(shared_xml)
        sheet_xml = input_zip.read(SHEET_PATH).decode("utf-8")

        for excel_row, result in results.items():
            decision = "APPROVE" if result.decision == "APPROVE" else "DECLINE"
            reference = f"{action_column}{excel_row}"
            self_closing = re.compile(
                rf'<c\b(?P<attrs>[^>]*\br="{re.escape(reference)}"[^>]*)\s*/>', re.S)
            normal = re.compile(
                rf'<c\b(?P<attrs>[^>]*\br="{re.escape(reference)}"[^>]*)>'
                rf'.*?</c>', re.S)

            def replace_cell(match):
                attrs = match.group("attrs").rstrip().removesuffix("/").rstrip()
                cell_type = "s" if has_shared_strings else "inlineStr"
                if re.search(r'\bt="[^"]*"', attrs):
                    attrs = re.sub(r'\bt="[^"]*"', f't="{cell_type}"', attrs)
                else:
                    attrs += f' t="{cell_type}"'
                body = (f"<v>{indexes[decision]}</v>" if has_shared_strings
                        else f"<is><t>{decision}</t></is>")
                return f"<c{attrs}>{body}</c>"

            sheet_xml, count = self_closing.subn(replace_cell, sheet_xml, count=1)
            if count == 0:
                sheet_xml, count = normal.subn(replace_cell, sheet_xml, count=1)
            if count != 1:
                raise ValueError(f"Could not update existing approval cell {reference}.")

            if result.corrected_amount_to_pay is not None and decision == "APPROVE":
                amount_reference = f"{amount_column}{excel_row}"
                amount_self_closing = re.compile(
                    rf'<c\b(?P<attrs>[^>]*\br="{re.escape(amount_reference)}"[^>]*)\s*/>', re.S)
                amount_normal = re.compile(
                    rf'<c\b(?P<attrs>[^>]*\br="{re.escape(amount_reference)}"[^>]*)>'
                    rf'.*?</c>', re.S)

                def replace_amount(match):
                    attrs = match.group("attrs").rstrip().removesuffix("/").rstrip()
                    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
                    value = f"{result.corrected_amount_to_pay:.2f}"
                    return f"<c{attrs}><v>{value}</v></c>"

                sheet_xml, amount_count = amount_self_closing.subn(
                    replace_amount, sheet_xml, count=1)
                if amount_count == 0:
                    sheet_xml, amount_count = amount_normal.subn(
                        replace_amount, sheet_xml, count=1)
                if amount_count != 1:
                    raise ValueError(
                        f"Could not update existing Finance Amount To Pay cell {amount_reference}.")

        with zipfile.ZipFile(destination, "w") as output_zip:
            for info in input_zip.infolist():
                data = input_zip.read(info.filename)
                if info.filename == SHARED_STRINGS_PATH and shared_xml is not None:
                    data = shared_xml.encode("utf-8")
                elif info.filename == SHEET_PATH:
                    data = sheet_xml.encode("utf-8")
                output_zip.writestr(info, data)
    return len(results)
