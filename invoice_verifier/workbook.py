from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .models import InvoiceRecord


REQUIRED_HEADERS = {
    "supplier": "Supplier Name",
    "brand": "Brand",
    "location": "Location",
    "order_number": "Order No.",
    "invoice_number": "Invoice No.",
    "currency": "PO Currency",
    "po_amount": "PO Amount",
    "invoice_date": "Invoice Date",
    "attachment_url": "File Attachment",
    "unique_reference": "Unique Reference",
    "record_id": "Record ID",
    "payment_status": "(FULL / PARTIAL) Payment Status",
    "amount_to_pay": "Finance Amount To Pay",
    "received_qty": "Finance Received Qty",
    "tax_code": "Finance Tax Code",
}


def _normalise_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def read_records(path: str | Path) -> list[InvoiceRecord]:
    """Read the source workbook without changing it."""
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    sheet = workbook.active
    # Some exported workbooks contain a stale worksheet dimension (A1:X1)
    # even though data continues below row 1. Reset it so read-only streaming
    # scans the real used range without modifying the source file.
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("The workbook is empty.") from exc

    indexed = {_normalise_header(value): index for index, value in enumerate(header_row)}
    missing = [label for label in REQUIRED_HEADERS.values() if label not in indexed]
    if missing:
        raise ValueError("Missing required column(s): " + ", ".join(missing))

    def get(row: tuple, key: str):
        return row[indexed[REQUIRED_HEADERS[key]]]

    records: list[InvoiceRecord] = []
    for excel_row, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        records.append(InvoiceRecord(
            excel_row=excel_row,
            supplier=_text(get(row, "supplier")),
            brand=_text(get(row, "brand")),
            location=_text(get(row, "location")),
            order_number=_text(get(row, "order_number")),
            invoice_number=_text(get(row, "invoice_number")),
            currency=_text(get(row, "currency")),
            po_amount=_number(get(row, "po_amount")),
            invoice_date=_text(get(row, "invoice_date")),
            attachment_url=_text(get(row, "attachment_url")),
            unique_reference=_text(get(row, "unique_reference")),
            record_id=_text(get(row, "record_id")),
            payment_status=_text(get(row, "payment_status")),
            amount_to_pay=_number(get(row, "amount_to_pay")),
            received_qty=_number(get(row, "received_qty")),
            tax_code=_text(get(row, "tax_code")),
            source={_normalise_header(header_row[i]): row[i] for i in range(len(header_row))},
        ))
    workbook.close()
    return records
